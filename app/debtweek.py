"""Еженедельный отчёт по дебиторке клиентов по филиалам → чат Bitrix24 «ANESCOM // DEBT».

Каждую пятницу в 19:00 (задача планировщика «ANESCOM debt», deploy/debt_send.py) портал строит Excel
по образцу ручного отчёта «ДЗ» (филиал → клиенты, долг, оплачено за неделю, статус), сохраняет
снимок долгов недели (таблица an_debt_weekly — для «изменения за неделю») и отправляет в чат
сводку текстом + файл. Данные — выгрузка 1С (an_debts, an_payments, an_sales).
Контакты клиентов в отчёт не попадают (правило: номера телефонов не пишем).

Настройки config.env: DEBT_CHAT_ID (ID чата, если не задан — ищем чат по названию DEBT_CHAT_TITLE).
"""
import base64
import logging
import os
from datetime import date, datetime, timedelta, timezone

import httpx
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response

from . import analytics, config, db

log = logging.getLogger("edo.debt")
router = APIRouter()

CHAT_TITLE = os.environ.get("DEBT_CHAT_TITLE", "ANESCOM // DEBT")
CHAT_ID = os.environ.get("DEBT_CHAT_ID", "").strip()
STALE_DAYS = 30          # «без оплат больше 30 дней» — красный статус


def short_branch(b: str) -> str:
    """«Филиал ТОО "ANESCOM/АНЕСКОМ" г: Атырау» → «Атырау»."""
    import re
    m = re.search(r"г[\.:]\s*([^,\"]+)", b or "")
    return m.group(1).strip() if m else (b or "Филиал не определён")


def _tz():
    return timezone(timedelta(hours=config.TZ_OFFSET_HOURS))


def _ensure(c):
    c.execute("""CREATE TABLE IF NOT EXISTS an_debt_weekly (
        week_end TEXT, client_key TEXT, client TEXT, branch TEXT, debt REAL, prepay REAL,
        PRIMARY KEY (week_end, client_key))""")


def build(day: date = None) -> dict:
    """Неделя: 7 дней по day включительно (по умолчанию — сегодня, пятница)."""
    from .dashboards import is_internal, client_branches, cname
    b = day or datetime.now(_tz()).date()
    a = b - timedelta(days=6)
    with analytics.conn() as c:
        _ensure(c)
        last = c.execute("SELECT MAX(ts) t FROM an_loads WHERE kind='debts'").fetchone()["t"]
        # филиал клиента — где он больше всего покупал (1С не присылает филиал у долгов);
        # в долгах ключ — контрагент, в продажах/оплатах — партнёр, поэтому сверяем и по названию
        cmap = client_branches(c)
        def by_name(sql, args=()):
            out = {}
            for r in c.execute(sql, args):
                out[cname(r["n"])] = r["v"] if cname(r["n"]) not in out else (
                    max(out[cname(r["n"])], r["v"]) if isinstance(r["v"], str) else out[cname(r["n"])] + (r["v"] or 0))
            return out
        paid_w = by_name("SELECT client n, SUM(sum) v FROM an_payments WHERE date BETWEEN ? AND ? GROUP BY n", (str(a), str(b)))
        last_pay = by_name("SELECT client n, MAX(date) v FROM an_payments GROUP BY n")
        ship_w = by_name("SELECT client n, SUM(sum) v FROM an_sales WHERE date BETWEEN ? AND ? GROUP BY n", (str(a), str(b)))
        prev_week = c.execute("SELECT MAX(week_end) w FROM an_debt_weekly WHERE week_end < ?", (str(b),)).fetchone()["w"]
        prev = {r["client_key"]: r["debt"] for r in c.execute(
            "SELECT client_key, debt FROM an_debt_weekly WHERE week_end=?", (prev_week,))} if prev_week else {}
        # долги клиента могут быть в нескольких организациях — складываем
        agg = {}
        for r in c.execute("SELECT client_key, client, manager, debt, prepay, overdue, overdue_days, branch FROM an_debts"):
            if is_internal(r["client"]):
                continue
            k = cname(r["client"]) or r["client_key"]
            x = agg.setdefault(k, {"key": k, "client": (r["client"] or "").strip(), "manager": r["manager"] or "",
                                   "debt": 0.0, "prepay": 0.0, "overdue": 0.0, "od": 0,
                                   "branch": short_branch(r["branch"] or (cmap.get(r["client_key"]) or cmap.get(k)
                                                                          or ("Филиал не определён",))[0])})
            x["debt"] += r["debt"] or 0
            x["prepay"] += r["prepay"] or 0
            x["overdue"] += r["overdue"] or 0
            x["od"] = max(x["od"], int(r["overdue_days"] or 0))
            if not x["manager"] and r["manager"]:
                x["manager"] = r["manager"]
    today = b
    clients = []
    for x in agg.values():
        net = x["debt"] - x["prepay"]
        x["net"] = net
        x["paid_w"] = paid_w.get(x["key"], 0) or 0
        x["ship_w"] = ship_w.get(x["key"], 0) or 0
        x["last_pay"] = last_pay.get(x["key"], "")
        try:
            x["days_np"] = (today - date.fromisoformat(x["last_pay"])).days if x["last_pay"] else None
        except ValueError:
            x["days_np"] = None
        x["delta"] = (x["debt"] - prev[x["key"]]) if (prev and x["key"] in prev) else (x["debt"] if prev else None)
        if x["debt"] <= 0:
            x["status"] = "ok"
        elif x["paid_w"] > 0:
            x["status"] = "paying"            # платит
        elif x["days_np"] is None or x["days_np"] > STALE_DAYS:
            x["status"] = "stale"             # нет оплат больше месяца / никогда
        else:
            x["status"] = "wait"              # на этой неделе не платил
        clients.append(x)
    debtors = sorted([x for x in clients if x["debt"] > 0.5], key=lambda x: (x["branch"], -x["debt"]))
    prepays = sorted([x for x in clients if x["prepay"] > 0.5], key=lambda x: -x["prepay"])
    branches = {}
    for x in debtors:
        br = branches.setdefault(x["branch"], {"br": x["branch"], "n": 0, "debt": 0.0, "overdue": 0.0, "paid_w": 0.0,
                                               "stale_n": 0, "stale": 0.0, "delta": 0.0, "ship_w": 0.0})
        br["n"] += 1
        br["debt"] += x["debt"]
        br["overdue"] += x["overdue"]
        br["paid_w"] += x["paid_w"]
        br["ship_w"] += x["ship_w"]
        br["delta"] += x["delta"] or 0
        if x["status"] == "stale":
            br["stale_n"] += 1
            br["stale"] += x["debt"]
    brs = sorted(branches.values(), key=lambda r: -r["debt"])
    tot = {k: sum(r[k] for r in brs) for k in ("n", "debt", "overdue", "paid_w", "stale_n", "stale", "delta", "ship_w")}
    paid_all = sum(v or 0 for v in paid_w.values())
    last_local = ""
    if last:
        try:
            last_local = datetime.fromisoformat(last).astimezone(_tz()).strftime("%d.%m.%Y %H:%M")
        except ValueError:
            last_local = last
    return {"a": a, "b": b, "debtors": debtors, "prepays": prepays, "branches": brs, "tot": tot,
            "prepay_total": sum(x["prepay"] for x in prepays), "paid_all": paid_all,
            "has_prev": bool(prev), "prev_week": prev_week, "last": last_local}


def save_snapshot(d: dict):
    with analytics.conn() as c:
        _ensure(c)
        c.execute("DELETE FROM an_debt_weekly WHERE week_end=?", (str(d["b"]),))
        c.executemany("INSERT INTO an_debt_weekly VALUES (?,?,?,?,?,?)",
                      [(str(d["b"]), x["key"], x["client"], x["branch"], x["debt"], x["prepay"])
                       for x in d["debtors"] + [p for p in d["prepays"] if p["debt"] <= 0.5]])
        c.commit()


# ------------------------------------------------------------------ Excel
STATUS = {"paying": ("платит", "E3F5EA", "0F5132"), "wait": ("нет оплаты на неделе", "FFF4D6", "7A5A00"),
          "stale": ("нет оплат > 30 дн.", "FDECEA", "B42318"), "ok": ("", None, None)}


def xlsx(d: dict, who: str = "EDO.ANESCOM") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import DataBarRule
    from . import xlsx_out as X
    wb = Workbook()
    per = f"{d['a']:%d.%m}–{d['b']:%d.%m.%Y}"
    sub = [f"Неделя {per} · данные 1С на {d['last'] or '—'}",
           "Без своих компаний. Статус: зелёный — платил на неделе, жёлтый — на неделе не платил, красный — нет оплат больше 30 дней."]

    # 1. Сводка по филиалам
    ws = wb.active
    ws.title = "Сводка"
    cols = [("Филиал", "text", 30), ("Клиентов с долгом", "int", 11), ("Дебиторка, ₸", "money", 17),
            ("Просрочено, ₸", "money", 15), ("Оплачено за неделю, ₸", "money", 16), ("Отгружено за неделю, ₸", "money", 16),
            ("Изменение долга за неделю, ₸", "money", 16), ("Без оплат > 30 дн.: клиентов", "int", 13),
            ("Без оплат > 30 дн., ₸", "money", 16), ("Доля филиала", "pct", 10)]
    r = X._header(ws, len(cols), "ANESCOM · дебиторская задолженность", f"Дебиторка клиентов по филиалам · {per}", sub)
    T = d["tot"]
    kp = [("Нам должны клиенты", T["debt"]), ("Оплачено за неделю (все клиенты)", d["paid_all"]),
          ("Без оплат больше 30 дней", T["stale"]), ("Переплаты клиентов (авансы)", d["prepay_total"])]
    for i, (lab, val) in enumerate(kp):
        cl = ws.cell(row=r, column=1 + i * 2 if i < 2 else 1 + i * 2, value=lab)
        cl.font = X.F_MUT
        v = ws.cell(row=r + 1, column=1 + i * 2, value=val)
        v.font = Font(name="Calibri", size=14, bold=True, color=X.RED if i in (0, 2) else X.NAVY)
        v.number_format = X.FMT["money"]
        v.alignment = Alignment(horizontal="left")
    r += 3
    rows = [[b_["br"], b_["n"], b_["debt"], b_["overdue"], b_["paid_w"], b_["ship_w"],
             b_["delta"] if d["has_prev"] else None, b_["stale_n"], b_["stale"],
             (b_["debt"] / T["debt"] * 100) if T["debt"] else 0] for b_ in d["branches"]]
    hr = r
    r = X._table(ws, r, cols, rows, total=["ИТОГО", T["n"], T["debt"], T["overdue"], T["paid_w"], T["ship_w"],
                                           T["delta"] if d["has_prev"] else None, T["stale_n"], T["stale"], 100.0],
                 key_col=3, freeze=False, autofilter=False)
    if rows:
        ws.conditional_formatting.add(f"J{hr + 1}:J{hr + len(rows)}",
                                      DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color="CC2B28"))
    if not d["has_prev"]:
        ws.cell(row=r, column=1, value="Изменение за неделю появится со следующего отчёта (сравнение со снимком прошлой пятницы).").font = X.F_MUT
        r += 1
    r += 1
    # топ-10 должников
    ws.cell(row=r, column=1, value="Топ-10 должников").font = X.F_SEC
    r += 1
    top = sorted(d["debtors"], key=lambda x: -x["debt"])[:10]
    tcols = [("Клиент", "text", 30), ("Филиал", "text", 11), ("Долг, ₸", "money", 17), ("Просрочено, ₸", "money", 15),
             ("Оплачено за неделю, ₸", "money", 16), ("Последняя оплата", "text", 16), ("Статус", "text", 16)]
    r0 = r
    r = X._table(ws, r, tcols, [[x["client"], x["branch"], x["debt"], x["overdue"], x["paid_w"],
                                 _d(x["last_pay"]), STATUS[x["status"]][0]] for x in top], freeze=False, autofilter=False)
    _paint_status(ws, r0 + 1, 7, top)
    X._foot(ws, r, [], who)
    X._print_setup(ws, True)

    # 2. По филиалам — как в ручном отчёте: блок филиала, клиенты, итог филиала
    ws = wb.create_sheet("По филиалам")
    cols = [("№", "int", 5), ("Клиент", "text", 40), ("Менеджер", "text", 24), ("Долг, ₸", "money", 16),
            ("Просрочено, ₸", "money", 14), ("Дней просрочки", "int", 9), ("Оплачено за неделю, ₸", "money", 15),
            ("Отгружено за неделю, ₸", "money", 15), ("Последняя оплата", "text", 13), ("Изменение за неделю, ₸", "money", 15),
            ("Статус", "text", 20)]
    r = X._header(ws, len(cols), "ANESCOM · дебиторская задолженность", f"Дебиторка по филиалам и клиентам · {per}", sub)
    head = r
    for j, (name, kind, w) in enumerate(cols, 1):
        cl = ws.cell(row=r, column=j, value=name)
        cl.font, cl.fill = X.F_HEAD, X.FILL_HEAD
        cl.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if kind == "text" else "right")
        ws.column_dimensions[cl.column_letter].width = w
    ws.row_dimensions[r].height = 32
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    r += 1
    band = PatternFill("solid", fgColor="0A2B52")
    for b_ in d["branches"]:
        for j in range(1, len(cols) + 1):
            ws.cell(row=r, column=j).fill = band
        cl = ws.cell(row=r, column=1, value=f"«АНЕСКОМ» — {b_['br'].upper()} · {_n(b_['n'])}")
        cl.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        s = ws.cell(row=r, column=4, value=b_["debt"])
        s.font, s.number_format = Font(name="Calibri", size=12, bold=True, color="FFFFFF"), X.FMT["money"]
        ws.row_dimensions[r].height = 22
        r += 1
        rows_b = [x for x in d["debtors"] if x["branch"] == b_["br"]]
        start = r
        for i, x in enumerate(rows_b, 1):
            vals = [i, x["client"], " ".join((x["manager"] or "").split()[:2]), x["debt"], x["overdue"] or None, x["od"] or None,
                    x["paid_w"] or None, x["ship_w"] or None, _d(x["last_pay"]) or "не было",
                    x["delta"] if d["has_prev"] else None, STATUS[x["status"]][0]]
            for j, ((name, kind, _w), v) in enumerate(zip(cols, vals), 1):
                cl = ws.cell(row=r, column=j, value=v)
                cl.font = X.F_CELL
                cl.border = Border(bottom=X.thin)
                if kind != "text":
                    cl.number_format = X.FMT[kind]
                if j == 4:
                    cl.font = Font(name="Calibri", size=11, bold=True, color=X.NAVY)
                if j == 10 and v:
                    cl.font = Font(name="Calibri", size=10.5, color="B42318" if v > 0 else "0F5132")
            _paint_status(ws, r, 11, [x])
            r += 1
        if rows_b:
            ws.conditional_formatting.add(f"D{start}:D{r - 1}", DataBarRule(start_type="min", end_type="max", color="F3B0AE"))
        tot = ["", f"Итого {b_['br']}", "", b_["debt"], b_["overdue"], None, b_["paid_w"], b_["ship_w"], "",
               b_["delta"] if d["has_prev"] else None, f"без оплат: {b_['stale_n']}"]
        for j, ((name, kind, _w), v) in enumerate(zip(cols, tot), 1):
            cl = ws.cell(row=r, column=j, value=v)
            cl.font, cl.fill = X.F_BOLD, X.FILL_TOT
            cl.border = Border(top=Side(style="medium", color=X.NAVY))
            if kind != "text":
                cl.number_format = X.FMT[kind]
        r += 2
    tot = ["", "ВСЕГО ПО КОМПАНИИ", "", T["debt"], T["overdue"], None, T["paid_w"], T["ship_w"], "",
           T["delta"] if d["has_prev"] else None, f"без оплат: {T['stale_n']}"]
    for j, ((name, kind, _w), v) in enumerate(zip(cols, tot), 1):
        cl = ws.cell(row=r, column=j, value=v)
        cl.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        cl.fill = PatternFill("solid", fgColor=X.RED)
        if kind != "text":
            cl.number_format = X.FMT[kind]
    X._foot(ws, r + 1, [], who)
    X._print_setup(ws, True, f"{head}:{head}")

    # 3. Переплаты
    ws = wb.create_sheet("Переплаты (авансы)")
    pcols = [("№", "int", 5), ("Клиент", "text", 42), ("Филиал", "text", 20), ("Менеджер", "text", 24),
             ("Переплата, ₸", "money", 16), ("Долг, ₸", "money", 14)]
    r = X._header(ws, len(pcols), "ANESCOM · авансы клиентов", f"Переплаты клиентов · на {d['b']:%d.%m.%Y}",
                  ["Деньги клиентов у нас: товар ещё не отгружен или не закрыт документом."])
    X._table(ws, r, pcols, [[i, x["client"], x["branch"], " ".join((x["manager"] or "").split()[:2]), x["prepay"], x["debt"] or None]
                            for i, x in enumerate(d["prepays"], 1)],
             total=["", "ИТОГО", "", "", d["prepay_total"], None], key_col=5)
    X._print_setup(ws, False)
    return X._bytes(wb)


def _n(n: int) -> str:
    n = int(n)
    w = "клиент" if n % 10 == 1 and n % 100 != 11 else ("клиента" if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14 else "клиентов")
    return f"{n} {w}"


def _d(s):
    try:
        return date.fromisoformat(s).strftime("%d.%m.%Y") if s else ""
    except ValueError:
        return s


def _paint_status(ws, r0, col, xs):
    from openpyxl.styles import Font, PatternFill
    for i, x in enumerate(xs):
        t, bg, fg = STATUS[x["status"]]
        if bg:
            cl = ws.cell(row=r0 + i, column=col)
            cl.fill = PatternFill("solid", fgColor=bg)
            cl.font = Font(name="Calibri", size=10, bold=True, color=fg)


# ------------------------------------------------------------------ текст в чат
def _m(v):
    v = float(v or 0)
    if abs(v) >= 1e9:
        return f"{v / 1e9:.2f} млрд".replace(".", ",")
    if abs(v) >= 1e6:
        return f"{v / 1e6:.1f} млн".replace(".", ",")
    return f"{v:,.0f}".replace(",", " ")


def text(d: dict) -> str:
    T = d["tot"]
    L = [f"[B]📊 Дебиторка клиентов · неделя {d['a']:%d.%m}–{d['b']:%d.%m.%Y}[/B]", ""]
    L.append(f"💰 Нам должны: [B]{_m(T['debt'])} ₸[/B] · {_n(T['n'])}"
             + (f" ({'▲ +' if T['delta'] > 0 else '▼ '}{_m(T['delta'])} за неделю)" if d["has_prev"] else ""))
    L.append(f"✅ Оплачено за неделю: [B]{_m(d['paid_all'])} ₸[/B]")
    L.append(f"🔴 Без оплат больше 30 дней: {_n(T['stale_n'])} на {_m(T['stale'])} ₸")
    L.append(f"💵 Авансы клиентов: {_m(d['prepay_total'])} ₸")
    L.append("")
    L.append("[B]По филиалам[/B]")
    for b_ in d["branches"]:
        L.append(f"• {b_['br']}: {_m(b_['debt'])} ₸ · оплачено {_m(b_['paid_w'])} · без оплат {b_['stale_n']}")
    top = sorted(d["debtors"], key=lambda x: -x["debt"])[:5]
    if top:
        L.append("")
        L.append("[B]Крупнейшие должники[/B]")
        icon = {"paying": "🟢", "wait": "🟡", "stale": "🔴", "ok": ""}
        for x in top:
            L.append(f"{icon[x['status']]} {x['client']} ({x['branch']}) — {_m(x['debt'])} ₸"
                     + (f", оплатил {_m(x['paid_w'])}" if x["paid_w"] else ""))
    L.append("")
    L.append(f"Подробно по клиентам — в файле Excel. Данные 1С на {d['last'] or '—'}.")
    return "\n".join(L)


# ------------------------------------------------------------------ отправка в Bitrix24
def _hook(method: str, params: dict):
    if not config.B24_WEBHOOK:
        raise RuntimeError("Не настроен вебхук Bitrix24 (EDO_B24_WEBHOOK)")
    r = httpx.post(f"{config.B24_WEBHOOK}/{method}.json", json=params, timeout=60)
    js = r.json()
    if "error" in js:
        raise RuntimeError(f"{method}: {js.get('error_description') or js['error']}")
    return js.get("result")


def find_chat() -> str:
    if CHAT_ID:
        return CHAT_ID if CHAT_ID.startswith("chat") else f"chat{CHAT_ID}"
    want = CHAT_TITLE.replace(" ", "").lower()
    for method, params in (("im.search.chat.list", {"FIND": "DEBT"}), ("im.recent.list", {})):
        try:
            res = _hook(method, params) or []
        except RuntimeError as e:
            log.warning("debt: %s: %s", method, e)
            continue
        items = res.get("items", res) if isinstance(res, dict) else res
        for it in items:
            ch = it.get("chat") if isinstance(it.get("chat"), dict) else it
            title = (ch.get("name") or ch.get("title") or it.get("title") or "")
            if title.replace(" ", "").lower() == want:
                return f"chat{ch.get('id')}"
    raise RuntimeError(f"Не найден чат «{CHAT_TITLE}» — впишите DEBT_CHAT_ID в config.env")


def send(day: date = None, dry: bool = False) -> dict:
    d = build(day)
    body = xlsx(d)
    name = f"Дебиторка_по_филиалам_{d['b']:%d.%m.%Y}.xlsx"
    out = {"file": name, "clients": len(d["debtors"])}
    if dry:
        return {**out, "text": text(d), "bytes": len(body)}
    dialog = find_chat()
    out["chat"] = dialog
    _hook("im.message.add", {"DIALOG_ID": dialog, "MESSAGE": text(d), "SYSTEM": "N"})
    try:
        chat_num = int(dialog.replace("chat", ""))
        folder = _hook("im.disk.folder.get", {"CHAT_ID": chat_num})
        up = _hook("disk.folder.uploadfile", {"id": folder["ID"], "data": {"NAME": name},
                                              "fileContent": [name, base64.b64encode(body).decode()],
                                              "generateUniqueName": True})
        _hook("im.disk.file.commit", {"CHAT_ID": chat_num, "FILE_ID": up["ID"]})
        out["upload"] = "ok"
    except Exception as e:      # нет права disk у вебхука — даём ссылку на портал
        log.warning("debt upload: %s", e)
        out["upload"] = f"ошибка: {e}"
        _hook("im.message.add", {"DIALOG_ID": dialog, "SYSTEM": "N",
                                 "MESSAGE": f"Файл Excel: {config.PUBLIC_BASE_URL}/dash/debtweek.xlsx (вход через Битрикс24)"})
    save_snapshot(d)
    with db.conn() as c:
        c.execute("INSERT INTO access_log (ts, b24_id, name, action, ip, ua, detail) VALUES (?,?,?,?,?,?,?)",
                  (db.now_iso(), None, "дебиторка", "debt_sent", "server", "", str(out)[:500]))
    return out


# ------------------------------------------------------------------ страницы
def _allowed(request: Request) -> bool:
    from . import access, main as m
    if not m.is_staff(request):
        return False
    s = access.scope(request)
    return s.password or s.role in ("admin", "director", "accountant", "rop")


@router.get("/dash/debtweek.xlsx")
def debt_xlsx(request: Request):
    from . import b24auth
    if not _allowed(request):
        raise HTTPException(403, "Отчёт по дебиторке — для бухгалтера, РОП, директора и администратора")
    d = build()
    name = f"debitorka_{d['b']:%Y%m%d}.xlsx"
    return Response(xlsx(d, b24auth.me(request)["name"]),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/dash/debtweek", response_class=HTMLResponse)
def debt_preview(request: Request, sent: str = ""):
    from . import b24auth, main as m
    import html as _h
    if not b24auth.is_admin(request):
        raise HTTPException(403, "Только для администратора")
    t = text(build())
    body = _h.escape(t).replace("[B]", "<b>").replace("[/B]", "</b>").replace("\n", "<br>")
    note = f"<p style='color:#0f5132'>Результат: {_h.escape(sent)}</p>" if sent else ""
    return HTMLResponse(f"""<!doctype html><meta charset=utf-8><meta name=viewport content="width=device-width,initial-scale=1">
<title>Дебиторка — предпросмотр</title><body style="font:15px/1.5 Roboto,Arial;background:#eef1f5;margin:0">
<div style="max-width:680px;margin:24px auto;background:#fff;padding:24px;border-radius:12px">
<p style="color:#6b7c90;margin-top:0">Так сообщение придёт в чат «{_h.escape(CHAT_TITLE)}» (каждую пятницу в 19:00) + файл Excel.</p>
{note}<div>{body}</div>
<p><a href="{m._root()}/dash/debtweek.xlsx">Скачать Excel</a></p>
<form method=post action="{m._root()}/dash/debtweek/send"><button style="background:#cc2b28;color:#fff;border:0;border-radius:999px;padding:10px 18px;font:500 14px Roboto">Отправить в чат сейчас</button></form>
</div></body>""")


@router.post("/dash/debtweek/send")
def debt_send_now(request: Request):
    from . import b24auth, main as m
    import urllib.parse
    if not b24auth.is_admin(request):
        raise HTTPException(403, "Только для администратора")
    try:
        res = send()
    except Exception as e:
        res = {"ошибка": str(e)}
    return RedirectResponse(f"{m._root()}/dash/debtweek?sent=" + urllib.parse.quote(str(res)), status_code=303)
