"""Excel-файлы портала: заказ поставщику, заявка на перемещение, отчёты вкладок.

Оформление одно на все файлы: шапка ANESCOM (логотип, название, дата, отбор), таблица с тёмно-синей
строкой заголовков, полосы, итоги, закреплённая шапка, автофильтр, печать на A4 по ширине листа.
"""
import io
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config

NAVY, RED, MUT, LINE, BAND = "001D3D", "CC2B28", "6B7C90", "DDE3EA", "F4F6F9"
_LOGO = Path(__file__).parent / "static" / "logo.png"
thin = Side(style="thin", color=LINE)
F_TITLE = Font(name="Calibri", size=16, bold=True, color=NAVY)
F_KICK = Font(name="Calibri", size=9, bold=True, color=RED)
F_MUT = Font(name="Calibri", size=10, color=MUT)
F_HEAD = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_CELL = Font(name="Calibri", size=10.5, color="1C2733")
F_BOLD = Font(name="Calibri", size=11, bold=True, color=NAVY)
F_SEC = Font(name="Calibri", size=12, bold=True, color=NAVY)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_BAND = PatternFill("solid", fgColor=BAND)
FILL_KEY = PatternFill("solid", fgColor="FDECEA")
FILL_TOT = PatternFill("solid", fgColor="E6ECF2")
FMT = {"money": '#,##0;[Red]-#,##0', "num": '#,##0.##', "int": '#,##0', "pct": '0.0"%"', "text": "@", "date": "DD.MM.YYYY"}


def now_local() -> datetime:
    return datetime.now(timezone(timedelta(hours=config.TZ_OFFSET_HOURS)))


def _header(ws, ncols: int, kick: str, title: str, lines: list) -> int:
    """Шапка листа. → номер первой свободной строки."""
    ws.sheet_view.showGridLines = False
    if _LOGO.exists():
        try:
            img = XLImage(str(_LOGO))
            img.height, img.width = 44, int(44 * 560 / 267)
            ws.add_image(img, "A1")
        except Exception:
            pass
    last = get_column_letter(max(ncols, 3))
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 14
    c = ws.cell(row=1, column=3 if ncols > 4 else 2, value=kick)
    c.font = F_KICK
    c = ws.cell(row=2, column=3 if ncols > 4 else 2, value=title)
    c.font = F_TITLE
    r = 4
    for ln in lines:
        ws.cell(row=r, column=1, value=ln).font = F_MUT
        r += 1
    for col in range(1, ncols + 1):
        ws.cell(row=r, column=col).border = Border(top=Side(style="medium", color=RED))
    ws.merge_cells(f"A{r}:{last}{r}") if ncols > 1 else None
    ws.row_dimensions[r].height = 6
    return r + 1


def _table(ws, r0: int, cols: list, rows: list, total=None, key_col=None, freeze=True, autofilter=True) -> int:
    """cols: [(заголовок, тип, ширина)]. → строка после таблицы."""
    for j, (name, kind, width) in enumerate(cols, 1):
        c = ws.cell(row=r0, column=j, value=name)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="right" if kind not in ("text",) else "left", vertical="center", wrap_text=True)
        cur = ws.column_dimensions[get_column_letter(j)].width or 0
        ws.column_dimensions[get_column_letter(j)].width = max(cur, width)
    ws.row_dimensions[r0].height = 30
    r = r0 + 1
    for i, row in enumerate(rows):
        for j, ((name, kind, _w), v) in enumerate(zip(cols, row), 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_CELL
            c.border = Border(bottom=thin)
            if kind != "text":
                c.number_format = FMT.get(kind, "General")
                c.alignment = Alignment(horizontal="right", vertical="top")
            else:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            if i % 2:
                c.fill = FILL_BAND
            if key_col and j == key_col:
                c.font, c.fill = Font(name="Calibri", size=11.5, bold=True, color=RED), FILL_KEY
        r += 1
    if freeze and rows:
        ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    if autofilter and rows:
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r - 1}"
    if total:
        for j, ((name, kind, _w), v) in enumerate(zip(cols, total), 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.fill = F_BOLD, FILL_TOT
            c.border = Border(top=Side(style="medium", color=NAVY))
            if kind != "text":
                c.number_format = FMT.get(kind, "General")
                c.alignment = Alignment(horizontal="right")
        r += 1
    return r


def _print_setup(ws, landscape=True, title_rows=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "EDO.ANESCOM · стр. &P из &N"
    if title_rows:
        ws.print_title_rows = title_rows


def _foot(ws, r: int, notes: list, who: str):
    r += 1
    for n in notes:
        ws.cell(row=r, column=1, value=n).font = F_MUT
        r += 1
    ws.cell(row=r, column=1, value=f"Сформировано {now_local():%d.%m.%Y %H:%M} · {who} · {config.PUBLIC_HOST}").font = F_MUT
    return r + 1


def _sign(ws, r: int, roles: list):
    r += 1
    for role in roles:
        ws.cell(row=r, column=1, value=role).font = F_CELL
        ws.cell(row=r, column=3, value="____________________  /  ____________________").font = F_MUT
        r += 2
    return r


def _bytes(wb) -> bytes:
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ------------------------------------------------------------------ заказ поставщику / перемещение
def order_book(data: dict, kind: str, org: str, wh: str, who: str) -> bytes:
    """kind=buy — заказ поставщику (+ лист перемещений), kind=move — заявка на перемещение."""
    wb = Workbook()
    ts = now_local()
    num_doc = ts.strftime("%Y%m%d-%H%M")
    sel = [f"Организация: {org or 'все'}", f"Склад: {wh or 'все склады'}",
           f"Остатки 1С на {data.get('anchor').strftime('%d.%m.%Y') if data.get('anchor') else '—'}"]
    reorder = data.get("reorder") or []

    def buy_sheet(ws):
        rows_src = sorted([r for r in reorder if r.get("buy", 0) > 0],
                          key=lambda r: (r["warehouse"], r.get("item_group") or "", r["item"]))
        cols = [("№", "int", 5), ("Номенклатура", "text", 56), ("Ед.", "text", 7), ("ЗАКАЗАТЬ", "num", 12),
                ("Остаток", "num", 10), ("Продано за 30 дн.", "num", 11), ("Хватит, дн.", "int", 9),
                ("Возьмём с других складов", "num", 13), ("Группа", "text", 20)]
        r = _header(ws, len(cols), "ANESCOM · заказ поставщику", f"Заказ поставщику № {num_doc} от {ts:%d.%m.%Y}", sel + [
            f"Позиций к заказу: {len(rows_src)}. Количество — до запаса на 30 дней по продажам за 30 дней, "
            "за вычетом того, что можно переместить с других складов (лист «Перемещения»)."])
        head_row = r
        _table(ws, r, cols, [], freeze=False, autofilter=False)
        ws.freeze_panes = ws.cell(row=r + 1, column=1)
        r += 1
        band = PatternFill("solid", fgColor="0A2B52")
        whs = []
        for x in rows_src:
            if x["warehouse"] not in whs:
                whs.append(x["warehouse"])
        n = 0
        for w in whs:
            items = [x for x in rows_src if x["warehouse"] == w]
            for j in range(1, len(cols) + 1):
                ws.cell(row=r, column=j).fill = band
            c = ws.cell(row=r, column=1, value=f"Склад получатель: {w} · позиций {len(items)}")
            c.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
            ws.row_dimensions[r].height = 22
            r += 1
            rows = []
            for x in items:
                n += 1
                rows.append([n, x["item"], x.get("unit") or "", float(x["buy"]), float(x["qty"] or 0),
                             float(x["sold_30"] or 0), round(x["days"] or 0), float(x.get("movable") or 0) or None,
                             x.get("item_group") or ""])
            # строки без повторной шапки
            for i, row in enumerate(rows):
                for j, ((name, kind, _w), v) in enumerate(zip(cols, row), 1):
                    c = ws.cell(row=r, column=j, value=v)
                    c.font, c.border = F_CELL, Border(bottom=thin)
                    if kind != "text":
                        c.number_format = FMT[kind]
                    if i % 2:
                        c.fill = FILL_BAND
                    if j == 4:
                        c.font, c.fill = Font(name="Calibri", size=12, bold=True, color=RED), FILL_KEY
                r += 1
            r += 1
        if not rows_src:
            ws.cell(row=r, column=1, value="Закупать нечего: всё закрывается перемещениями или запаса хватает.").font = F_MUT
            r += 1
        r = _sign(ws, r, ["Составил (закупщик)", "Согласовал (руководитель)"])
        _foot(ws, r, ["Цены и поставщика укажите при отправке; файл можно править."], who)
        _print_setup(ws, True, f"{head_row}:{head_row}")

    def move_sheet(ws, title):
        mv = sorted(data.get("move_rows") or [], key=lambda m: (m["src"], m["to"], m["item"]))
        cols = [("№", "int", 5), ("Откуда (склад отправитель)", "text", 26), ("Куда (склад получатель)", "text", 26),
                ("Номенклатура", "text", 50), ("Ед.", "text", 7), ("ПЕРЕМЕСТИТЬ", "num", 13), ("Группа", "text", 22)]
        r = _header(ws, len(cols), "ANESCOM · перемещение между складами", title, sel + [
            f"Строк: {len(mv)}. Берём только лишнее — сверх 30 дней продаж склада-отправителя."])
        rows = [[i, m["src"], m["to"], m["item"], m.get("unit") or "", float(m["qty"]), m.get("group") or ""]
                for i, m in enumerate(mv, 1)]
        head_row = r
        r = _table(ws, r, cols, rows, key_col=6)
        if not rows:
            ws.cell(row=r, column=1, value="Перемещать нечего: на других складах нет лишнего.").font = F_MUT
            r += 1
        r = _sign(ws, r, ["Отпустил (склад отправитель)", "Принял (склад получатель)"])
        _foot(ws, r, [], who)
        _print_setup(ws, True, f"{head_row}:{head_row}")

    if kind == "buy":
        ws = wb.active
        ws.title = "Заказ поставщику"
        buy_sheet(ws)
        move_sheet(wb.create_sheet("Перемещения"), f"Перемещения перед закупкой от {ts:%d.%m.%Y}")
    else:
        ws = wb.active
        ws.title = "Перемещение"
        move_sheet(ws, f"Заявка на перемещение № {num_doc} от {ts:%d.%m.%Y}")
    return _bytes(wb)


# ------------------------------------------------------------------ отчёт вкладки
def report_book(rep: dict, who: str) -> bytes:
    """rep — см. reports.build(): title, kick, sub, kpis, tables."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    width = max([len(t["cols"]) for t in rep["tables"]] + [4])
    r = _header(ws, width, rep["kick"], rep["title"], rep["sub"])
    # показатели — две колонки «название | значение»
    if rep["kpis"]:
        for label, val, _cls in rep["kpis"]:
            ws.cell(row=r, column=1, value=label).font = F_MUT
            c = ws.cell(row=r, column=2, value=val)
            c.font = F_BOLD
            if isinstance(val, (int, float)):
                c.number_format = FMT["money"]
                c.alignment = Alignment(horizontal="left")
            r += 1
        r += 1
    for t in rep["tables"]:
        ws.cell(row=r, column=1, value=t["title"]).font = F_SEC
        r += 1
        cols = [(name, kind, w) for name, kind, w in t["cols"]]
        r = _table(ws, r, cols, t["rows"], total=t.get("total"), freeze=False, autofilter=False)
        if t.get("note"):
            ws.cell(row=r, column=1, value=t["note"]).font = F_MUT
            r += 1
        r += 1
    _foot(ws, r, rep.get("notes") or [], who)
    _print_setup(ws, width > 5)
    # каждая таблица ещё и отдельным листом — удобно фильтровать
    for t in rep["tables"]:
        if len(t["rows"]) < 8:
            continue
        name = "".join(ch for ch in t["title"] if ch not in '[]:*?/\\')[:31] or "Лист"
        w2 = wb.create_sheet(name)
        r2 = _header(w2, len(t["cols"]), rep["kick"], t["title"], rep["sub"][:2])
        _table(w2, r2, t["cols"], t["rows"], total=t.get("total"))
        _print_setup(w2, len(t["cols"]) > 5)
    return _bytes(wb)
