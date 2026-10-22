"""Вкладка «Товары и группы» (22.10.2026): анализ товара / группы / бренда / поставщика — как отчёт GEZE.

/dash/items          — форма отбора и просмотр
/dash/items.pdf      — тот же отчёт в PDF (reportlab, A4 альбомная)
/dash/items.xlsx     — тот же отчёт в Excel (листы: Итоги, По группам, По товарам, Неликвид)
/dash/items/options  — подсказки для поля отбора (группы, поставщики, товары)

Отбор (mode): group — группа номенклатуры 1С; name — слово в наименовании (бренд, например GEZE);
item — один товар (item_key); supplier — всё, что покупали у поставщика.
Данные (analytics.sqlite3):
  an_cost_month — остатки и приход по себестоимости по месяцам (стоимость поступления + доп. расходы: доставка, пошлина);
  an_purchases  — закупки у поставщиков (поставщик, цена поставщика без НДС);
  an_sales      — продажи клиентам со скидками и НДС (продажи своим компаниям — GREENTE, АВТОКОМП — не считаются);
  an_stock      — текущие остатки.
Выручка в 1С — с НДС; прибыль считается без НДС (с 2026 г. — 16%, раньше — 12%), если 1С не прислала сумму НДС.
Доступ: все роли аналитики (решение пользователя 22.10.2026 16:10), в т.ч. менеджеры.
"""
import io
from calendar import monthrange
from datetime import date, timedelta
from typing import Optional

from fastapi import APIRouter, Request
from fastapi.responses import HTMLResponse, JSONResponse, Response

from . import analytics, config

router = APIRouter()

MODES = {"group": "Группа товара", "name": "Название / бренд", "item": "Один товар", "supplier": "Поставщик"}
ROLES = {"admin", "director", "rop", "accountant", "manager", "storekeeper", "purchaser"}


# ------------------------------------------------------------------ помощники
def _d(s: str, default: date) -> date:
    try:
        return date.fromisoformat((s or "").strip()[:10])
    except ValueError:
        return default


def _month_end(d: date) -> date:
    return date(d.year, d.month, monthrange(d.year, d.month)[1])


def _vat_rate(iso: str) -> float:
    return 0.16 if iso >= "2026-01-01" else 0.12


def _internal(name: str) -> bool:
    n = (name or "").upper()
    return any(k in n for k in config.INTERNAL_CLIENTS)


def _has(c, t: str) -> bool:
    return bool(c.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (t,)).fetchone())


def _cols(c, t: str) -> set:
    return {r[1] for r in c.execute(f"PRAGMA table_info({t})")}


def money(v) -> str:
    return f"{(v or 0):,.0f}".replace(",", " ")


def mln(v) -> str:
    return f"{(v or 0) / 1e6:,.2f}".replace(",", " ").replace(".", ",")


def _months(v) -> str:
    return "больше 5 лет" if v > 60 else f"≈{v:.0f} мес."


def pct(v, digits=0) -> str:
    return "—" if v is None else (f"{v:.{digits}f}%").replace(".", ",")


# ------------------------------------------------------------------ подсказки для отбора
def options(mode: str, q: str = "", limit: int = 40) -> list:
    q = (q or "").strip().lower()
    with analytics.conn() as c:
        if mode == "group":
            rows = [r[0] for r in c.execute(
                "SELECT item_group FROM an_sales WHERE item_group<>'' GROUP BY item_group "
                "UNION SELECT item_group FROM an_stock WHERE item_group<>'' GROUP BY item_group")]
            vals = [(g, g) for g in rows]
        elif mode == "supplier":
            vals = [(r[0], r[0]) for r in c.execute(
                "SELECT supplier FROM an_purchases WHERE supplier<>'' GROUP BY supplier ORDER BY SUM(sum) DESC")] \
                if _has(c, "an_purchases") else []
        elif mode == "item":
            vals = [(r[1], r[0]) for r in c.execute(
                "SELECT item, item_key FROM an_stock WHERE item_key<>'' GROUP BY item_key "
                "UNION SELECT item, item_key FROM an_sales WHERE item_key<>'' AND date>=? GROUP BY item_key",
                ((date.today() - timedelta(days=730)).isoformat(),))]
        else:
            return []
    if q:
        vals = [v for v in vals if q in v[1].lower()]
    vals.sort(key=lambda v: v[1].lower())
    return [{"value": v, "label": l} for v, l in vals[:limit]]


# ------------------------------------------------------------------ расчёт
def _keys(c, mode: str, q: str) -> Optional[set]:
    """item_key выбранных товаров (None — нечего искать)."""
    q = (q or "").strip()
    if not q:
        return None
    tables = [t for t in ("an_sales", "an_stock", "an_cost_month", "an_purchases") if _has(c, t)]
    keys = set()
    if mode == "item":
        return {q}
    if mode == "supplier":
        if not _has(c, "an_purchases"):
            return set()
        return {r[0] for r in c.execute("SELECT DISTINCT item_key FROM an_purchases WHERE supplier=?", (q,))}
    ql = q.lower()
    for t in tables:
        for item, key, grp in c.execute(f"SELECT item, item_key, item_group FROM {t} GROUP BY item_key"):
            if mode == "group" and (grp or "") == q:
                keys.add(key)
            elif mode == "name" and ql in (item or "").lower():
                keys.add(key)
    return keys


def _in(keys) -> tuple:
    ks = sorted(keys)
    return "(" + ",".join("?" * len(ks)) + ")", ks


def build(mode: str, q: str, a: date, b: date, s_date: date, org: str = "") -> dict:
    mode = mode if mode in MODES else "name"
    res = {"mode": mode, "mode_title": MODES[mode], "q": q, "a": a, "b": b, "org": org, "empty": True,
           "items": [], "groups": [], "T": {}, "notes": [], "recs": [], "has_disc": False, "has_extra": False}
    with analytics.conn() as c:
        keys = _keys(c, mode, q)
        if keys is None:
            res["why"] = "Выберите, по чему строить отчёт."
            return res
        if not keys:
            res["why"] = "По этому отбору в выгрузке 1С ничего не нашлось."
            return res
        ph, ks = _in(keys)
        orgw, orga = (" AND org=?", [org]) if org else ("", [])
        it = {}

        def row(key, item="", grp=""):
            r = it.get(key)
            if r is None:
                r = it[key] = {"key": key, "item": item, "group": grp, "q0": 0.0, "c0": 0.0, "qin": 0.0, "cin": 0.0,
                               "extra": 0.0, "pq": 0.0, "psum": 0.0, "suppliers": set(), "qs": 0.0, "rev": 0.0,
                               "disc": 0.0, "vat": 0.0, "cost": 0.0, "cost_est": 0.0, "q1": 0.0, "c1": 0.0,
                               "last_sale": ""}
            if item and not r["item"]:
                r["item"] = item
            if grp and not r["group"]:
                r["group"] = grp
            return r

        # --- остаток на дату и приход: по месяцам себестоимости
        m_last = _month_end(s_date)
        if m_last != s_date:  # не конец месяца → берём конец предыдущего месяца
            m_last = s_date.replace(day=1) - timedelta(days=1)
        res["s_date"] = m_last
        cm = _has(c, "an_cost_month")
        res["has_cost_month"] = cm and bool(c.execute("SELECT 1 FROM an_cost_month LIMIT 1").fetchone())
        if cm:
            # остаток на конец месяца m_last = qty_end последней строки товара с месяцем ≤ m_last
            for r in c.execute(
                    f"SELECT item_key, item, item_group, SUM(qty_end) q, SUM(cost_end) s FROM an_cost_month t "
                    f"WHERE item_key IN {ph}{orgw} AND date = (SELECT MAX(date) FROM an_cost_month x WHERE x.item_key=t.item_key "
                    f"AND x.org=t.org AND x.date<=?) GROUP BY item_key", ks + orga + [m_last.replace(day=1).isoformat()]):
                x = row(r["item_key"], r["item"], r["item_group"])
                x["q0"] += r["q"] or 0
                x["c0"] += r["s"] or 0
            for r in c.execute(
                    f"SELECT item_key, item, item_group, SUM(qty_in) q, SUM(cost_in) s, SUM(extra_in) e FROM an_cost_month "
                    f"WHERE item_key IN {ph}{orgw} AND date BETWEEN ? AND ? GROUP BY item_key",
                    ks + orga + [a.replace(day=1).isoformat(), b.isoformat()]):
                x = row(r["item_key"], r["item"], r["item_group"])
                x["qin"] += r["q"] or 0
                x["cin"] += (r["s"] or 0) + (r["e"] or 0)
                x["extra"] += r["e"] or 0
        # --- закупки у поставщиков (цена поставщика, поставщики)
        if _has(c, "an_purchases"):
            for r in c.execute(
                    f"SELECT item_key, item, item_group, supplier, SUM(qty) q, SUM(sum) s FROM an_purchases "
                    f"WHERE item_key IN {ph}{orgw} AND date BETWEEN ? AND ? GROUP BY item_key, supplier",
                    ks + orga + [a.isoformat(), b.isoformat()]):
                x = row(r["item_key"], r["item"], r["item_group"])
                x["pq"] += r["q"] or 0
                x["psum"] += r["s"] or 0
                if r["supplier"]:
                    x["suppliers"].add(r["supplier"])
        # --- продажи клиентам
        sc = _cols(c, "an_sales")
        dsel = ", SUM(disc_auto + disc_manual) d, SUM(vat) v" if "disc_auto" in sc else ", 0 d, 0 v"
        for r in c.execute(
                f"SELECT item_key, item, item_group, client, date, SUM(qty) q, SUM(sum) s, SUM(cost) cs{dsel} "
                f"FROM an_sales WHERE item_key IN {ph}{orgw} AND date BETWEEN ? AND ? "
                f"GROUP BY item_key, client, date", ks + orga + [a.isoformat(), b.isoformat()]):
            if _internal(r["client"]):
                continue
            x = row(r["item_key"], r["item"], r["item_group"])
            x["qs"] += r["q"] or 0
            x["rev"] += r["s"] or 0
            x["disc"] += r["d"] or 0
            vat = r["v"] or 0
            if not vat and r["s"]:
                rate = _vat_rate(r["date"])
                vat = r["s"] * rate / (1 + rate)
            x["vat"] += vat
            x["cost"] += r["cs"] or 0
            if r["date"] > x["last_sale"]:
                x["last_sale"] = r["date"]
        # --- текущий остаток
        for r in c.execute(
                f"SELECT item_key, item, item_group, SUM(qty) q, SUM(cost_sum) s FROM an_stock "
                f"WHERE item_key IN {ph} GROUP BY item_key", ks):
            x = row(r["item_key"], r["item"], r["item_group"])
            x["q1"] += r["q"] or 0
            x["c1"] += r["s"] or 0
        last = c.execute("SELECT MAX(ts) FROM an_loads WHERE kind IN ('sales','stock')").fetchone()[0]
    res["last"] = last

    months = max(1.0, ((b - a).days + 1) / 30.44)
    res["months"] = months
    items = []
    for x in it.values():
        if not any(x[k] for k in ("q0", "qin", "qs", "q1", "pq")):
            continue
        # себестоимость продаж: если 1С ещё не посчитала (месяц не закрыт) — оценка по средней цене единицы
        unit = (x["c1"] / x["q1"]) if x["q1"] > 0 and x["c1"] > 0 else \
               (x["cin"] / x["qin"]) if x["qin"] > 0 else (x["c0"] / x["q0"]) if x["q0"] > 0 else 0
        if x["qs"] > 0 and x["cost"] <= 0 and unit:
            x["cost_est"] = unit * x["qs"]
        cost = x["cost"] if x["cost"] > 0 else x["cost_est"]
        x["cost_used"] = cost
        x["rev_nv"] = x["rev"] - x["vat"]
        x["margin"] = x["rev_nv"] - cost if x["rev"] else 0.0
        x["mpct"] = (x["margin"] / x["rev_nv"] * 100) if x["rev_nv"] > 0 else None
        x["dpct"] = (x["disc"] / (x["rev"] + x["disc"]) * 100) if (x["rev"] + x["disc"]) > 0 and x["disc"] else None
        x["avg_price"] = x["rev"] / x["qs"] if x["qs"] else None
        x["ucost"] = cost / x["qs"] if x["qs"] else None
        x["u_in"] = x["cin"] / x["qin"] if x["qin"] else None
        x["u_sup"] = x["psum"] / x["pq"] if x["pq"] else None
        x["extra_pct"] = (x["extra"] / (x["cin"] - x["extra"]) * 100) if x["cin"] - x["extra"] > 0 and x["extra"] else None
        x["stock_months"] = (x["q1"] / (x["qs"] / months)) if x["qs"] > 0 and x["q1"] > 0 else None
        x["dead"] = x["q1"] > 0 and x["qs"] <= 0
        x["suppliers"] = ", ".join(sorted(x["suppliers"]))
        items.append(x)
    items.sort(key=lambda x: (-x["rev"], -x["c1"]))
    res["items"] = items
    if not items:
        res["why"] = "По этому отбору за период нет ни остатков, ни движения."
        return res
    res["empty"] = False

    def total(rows):
        T = {k: sum(r[k] for r in rows) for k in ("q0", "c0", "qin", "cin", "extra", "pq", "psum", "qs", "rev",
                                                  "disc", "vat", "cost_used", "cost_est", "q1", "c1", "rev_nv", "margin")}
        T["mpct"] = T["margin"] / T["rev_nv"] * 100 if T["rev_nv"] > 0 else None
        T["dpct"] = T["disc"] / (T["rev"] + T["disc"]) * 100 if T["disc"] else None
        T["extra_pct"] = T["extra"] / (T["cin"] - T["extra"]) * 100 if T["cin"] - T["extra"] > 0 and T["extra"] else None
        T["stock_months"] = T["q1"] / (T["qs"] / months) if T["qs"] > 0 else None
        T["n"] = len(rows)
        T["dead_q"] = sum(r["q1"] for r in rows if r["dead"])
        T["dead_c"] = sum(r["c1"] for r in rows if r["dead"])
        T["dead_n"] = sum(1 for r in rows if r["dead"])
        return T

    T = res["T"] = total(items)
    groups = {}
    for x in items:
        groups.setdefault(x["group"] or "(без группы)", []).append(x)
    res["groups"] = sorted(({"group": g, **total(rows)} for g, rows in groups.items()), key=lambda g: -g["rev"])
    res["has_disc"] = T["disc"] > 0
    res["has_extra"] = T["extra"] > 0
    res["recs"], res["notes"] = _advice(res)
    return res


def _advice(R) -> tuple:
    T, items = R["T"], R["items"]
    recs, notes = [], []
    if T["qin"] and T["qs"] and T["qin"] > 1.3 * T["qs"]:
        recs.append(f"Закупили в {str(round(T['qin'] / T['qs'], 1)).replace('.', ',')} раза больше, чем продали клиентам: пришло {money(T['qin'])}, "
                    f"продано {money(T['qs'])}. Остаток по себестоимости: {mln(T['c0'])} → {mln(T['c1'])} млн ₸.")
    if T["mpct"] is not None:
        line = f"Валовая прибыль без НДС — {mln(T['margin'])} млн ₸, {pct(T['mpct'], 1)} от выручки без НДС"
        if T["dpct"]:
            line += f"; скидками отдано {mln(T['disc'])} млн ₸ ({pct(T['dpct'])} от прайса)"
        recs.append(line + ".")
    neg = sorted([x for x in items if x["mpct"] is not None and x["mpct"] < 0 and x["qs"] > 0], key=lambda x: x["margin"])
    if neg:
        lst = "; ".join(f"{x['item']} ({pct(x['mpct'])})" for x in neg[:5])
        recs.append(f"Продаются в минус: {lst}. Чтобы получить 20% маржи без НДС, средняя цена продажи с НДС должна быть "
                    f"не ниже себестоимость × 1,16 ÷ 0,8 — поднять прайс или ограничить скидку на эти позиции.")
    low = [x for x in items if x["mpct"] is not None and 0 <= x["mpct"] < 10 and x["qs"] > 0]
    if low:
        recs.append("Маржа ниже 10%: " + "; ".join(f"{x['item']} ({pct(x['mpct'])})" for x in low[:5]) + ".")
    if T["dpct"] and T["dpct"] >= 20 and T["mpct"] is not None and T["mpct"] < 25:
        recs.append(f"Скидка {pct(T['dpct'])} при марже {pct(T['mpct'])}: скидку на позиции с маржой ниже 20% "
                    f"ограничить 10–15%, крупную скидку давать только под объект.")
    over = sorted([x for x in items if x["stock_months"] and x["stock_months"] > 12], key=lambda x: -x["c1"])
    if over:
        recs.append("Запас больше чем на год — остановить закупку, пока не станет меньше 6 месяцев: "
                    + "; ".join(f"{x['item']} — {money(x['q1'])} шт, {_months(x['stock_months'])}" for x in over[:5]) + ".")
    if T["dead_n"]:
        recs.append(f"Неликвид: {T['dead_n']} позиций, {money(T['dead_q'])} шт на {mln(T['dead_c'])} млн ₸ — за период "
                    f"ни одной продажи. Предложить в комплектах, объектах или распродать без ухода в минус.")
    if T["cost_est"]:
        notes.append(f"Себестоимость части продаж ({mln(T['cost_est'])} млн ₸) — оценка по средней цене единицы: "
                     "в 1С месяц ещё не закрыт (расчёт себестоимости не выполнен).")
    if not R["has_disc"]:
        notes.append("Скидки в выгрузке 1С появятся после обновления расширения «ANESCOM: Аналитика» (22.10.2026).")
    if not R.get("has_cost_month"):
        notes.append("Остаток на дату и закупка с доставкой и пошлиной появятся после обновления расширения 1С "
                     "(выгрузка себестоимости по месяцам).")
    notes.append("Выручка в 1С — с НДС; прибыль посчитана без НДС (2026 — 16%, раньше — 12%). "
                 "Продажи своим компаниям (GREENTE, АВТОКОМП) не входят в продажи клиентам.")
    return recs, notes


# ------------------------------------------------------------------ доступ и параметры
def _gate(request: Request):
    from . import main as m, access
    if not m.is_staff(request):
        return m.templates.TemplateResponse("login.html", {
            "request": request, "brand": config.BRAND, "company": config.COMPANY, "base": m._root(), "b24": True})
    s = access.scope(request)
    if s.password or s.full_dash or s.role in ROLES:
        return None
    return m.templates.TemplateResponse("noaccess.html", {
        "request": request, "brand": config.BRAND, "base": m._root(), "title": "Нет доступа",
        "text": "Отчёты по товарам и группам — для сотрудников с ролью в аналитике."})


def _params(qp) -> dict:
    t = date.today()
    a = _d(qp.get("a", ""), date(t.year, 1, 1))
    b = _d(qp.get("b", ""), t)
    if a > b:
        a, b = b, a
    s = _d(qp.get("s", ""), a - timedelta(days=1))
    return {"mode": qp.get("mode", "name") if qp.get("mode") in MODES else "name", "q": (qp.get("q") or "").strip(),
            "a": a, "b": b, "s": s, "org": (qp.get("org") or "").strip()}


def _title(R) -> str:
    q = R["q"]
    if R["mode"] == "item" and R["items"]:
        q = R["items"][0]["item"]
    return f"{R['mode_title']}: {q}"


def _disp(R, ext, inline=False) -> str:
    """Content-Disposition: латиница в filename, полное имя (с кириллицей) — в filename*."""
    import re
    import urllib.parse as up
    q = re.sub(r"[^0-9A-Za-zА-Яа-яЁё]+", "_", R["q"])[:40].strip("_") or "otchet"
    full = f"tovary_{q}_{R['b']:%Y%m%d}.{ext}"
    ascii_ = re.sub(r"[^0-9A-Za-z_.-]", "", full.replace("tovary__", "tovary_")) or f"tovary.{ext}"
    return f'{"inline" if inline else "attachment"}; filename="{ascii_}"; filename*=UTF-8\'\'{up.quote(full)}'


@router.get("/dash/items/options")
def items_options(request: Request, mode: str = "name", q: str = ""):
    if _gate(request) is not None:
        return JSONResponse([], status_code=401)
    return options(mode, q)


@router.get("/dash/items", response_class=HTMLResponse)
def dash_items(request: Request):
    from . import main as m, dashboards as d
    g = _gate(request)
    if g is not None:
        return g
    P = _params(request.query_params)
    R = build(P["mode"], P["q"], P["a"], P["b"], P["s"], P["org"]) if P["q"] else \
        {"empty": True, "why": "", "mode": P["mode"], "q": ""}
    import urllib.parse as up
    rq = up.urlencode({"mode": P["mode"], "q": P["q"], "a": P["a"].isoformat(), "b": P["b"].isoformat(),
                       "s": P["s"].isoformat(), "org": P["org"]})
    data = {"empty": False}
    return d.templates.TemplateResponse("dash_items.html", {
        "request": request, "brand": config.BRAND, "base": m._root(), "p": "month", "org": "", "branch": "", "d": "",
        "periods": d.periods_all(), "period_title": "", "nav": d.NAV, "tab": "items", "page_title": "Товары и группы",
        "show_branch": False, "data": data, "R": R, "P": P, "MODES": MODES, "rq": rq, "money": money, "mln": mln,
        "pct": pct, "title": _title(R) if not R["empty"] else "", "human_date": d.human_date, "last_local": "",
    })


@router.get("/dash/items.xlsx")
def dash_items_xlsx(request: Request):
    from . import b24auth
    g = _gate(request)
    if g is not None:
        return g
    P = _params(request.query_params)
    R = build(P["mode"], P["q"], P["a"], P["b"], P["s"], P["org"])
    if R["empty"]:
        return HTMLResponse(R.get("why") or "Нет данных", status_code=404)
    who = b24auth.me(request)["name"]
    b24auth.log_access(request, "report", b24auth.current_user(request), f"товары xlsx {P['mode']} {P['q']}")
    return Response(xlsx(R, who), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": _disp(R, "xlsx")})


@router.get("/dash/items.pdf")
def dash_items_pdf(request: Request):
    from . import b24auth
    g = _gate(request)
    if g is not None:
        return g
    P = _params(request.query_params)
    R = build(P["mode"], P["q"], P["a"], P["b"], P["s"], P["org"])
    if R["empty"]:
        return HTMLResponse(R.get("why") or "Нет данных", status_code=404)
    who = b24auth.me(request)["name"]
    b24auth.log_access(request, "report", b24auth.current_user(request), f"товары pdf {P['mode']} {P['q']}")
    return Response(pdf(R, who), media_type="application/pdf",
                    headers={"Content-Disposition": _disp(R, "pdf", inline=True)})


# ------------------------------------------------------------------ таблицы (общие для PDF и Excel)
def item_table(R, with_group=True):
    head = ["Товар", "Группа", f"Остаток {R['s_date']:%d.%m.%y}, шт", "Остаток, ₸", "Пришло, шт", "Цена пост., ₸/шт",
            "С дост. и пошл., ₸/шт", "Доп. расх.", "Продано, шт", "Выручка с НДС, ₸", "Скидка", "Ср. цена с НДС, ₸",
            "Себест., ₸/шт", "Прибыль без НДС, ₸", "Маржа", "Остаток сейчас, шт", "Остаток сейчас, ₸", "Запас, мес."]
    kinds = ["text", "text", "int", "money", "int", "money", "money", "pct", "int", "money", "pct", "money", "money",
             "money", "pct", "int", "money", "num"]
    rows = []
    for x in R["items"]:
        rows.append([x["item"], x["group"], x["q0"], x["c0"], x["qin"], x["u_sup"], x["u_in"], x["extra_pct"], x["qs"],
                     x["rev"], x["dpct"], x["avg_price"], x["ucost"], x["margin"] if x["qs"] else None, x["mpct"],
                     x["q1"], x["c1"], ("нет продаж" if x["dead"] else x["stock_months"])])
    T = R["T"]
    tot = ["Итого", "", T["q0"], T["c0"], T["qin"], None, None, T["extra_pct"], T["qs"], T["rev"], T["dpct"], None, None,
           T["margin"], T["mpct"], T["q1"], T["c1"], T["stock_months"]]
    if not with_group:
        head, kinds, rows, tot = head[:1] + head[2:], kinds[:1] + kinds[2:], [r[:1] + r[2:] for r in rows], tot[:1] + tot[2:]
    return head, kinds, rows, tot


def group_table(R):
    head = ["Группа", "Позиций", f"Остаток {R['s_date']:%d.%m.%y}, ₸", "Приход (с доп. расх.), ₸", "Продано, шт",
            "Выручка с НДС, ₸", "Скидка", "Прибыль без НДС, ₸", "Маржа", "Остаток сейчас, ₸", "Запас, мес.", "Неликвид, ₸"]
    kinds = ["text", "int", "money", "money", "int", "money", "pct", "money", "pct", "money", "num", "money"]
    rows = [[g["group"], g["n"], g["c0"], g["cin"], g["qs"], g["rev"], g["dpct"], g["margin"], g["mpct"], g["c1"],
             g["stock_months"], g["dead_c"]] for g in R["groups"]]
    T = R["T"]
    tot = ["Итого", T["n"], T["c0"], T["cin"], T["qs"], T["rev"], T["dpct"], T["margin"], T["mpct"], T["c1"],
           T["stock_months"], T["dead_c"]]
    return head, kinds, rows, tot


def kpis(R) -> list:
    T = R["T"]
    k = [(f"{money(T['q0'])} шт · {mln(T['c0'])} млн ₸", f"остаток на {R['s_date']:%d.%m.%Y} по себестоимости без НДС"),
         (f"{money(T['qin'])} шт · {mln(T['cin'])} млн ₸",
          "пришло за период с доставкой и пошлиной" + (f"; доп. расходы +{pct(T['extra_pct'], 1)}" if T["extra_pct"] else "")),
         (f"{money(T['qs'])} шт · {mln(T['rev'])} млн ₸", f"продано клиентам, выручка с НДС после скидок; без НДС {mln(T['rev_nv'])} млн"),
         (f"{pct(T['dpct'])} · {mln(T['disc'])} млн ₸" if T["disc"] else "—", "отдано скидками (авто + ручные)"),
         (f"{pct(T['mpct'], 1)} · {mln(T['margin'])} млн ₸", "валовая прибыль без НДС"),
         (f"{money(T['q1'])} шт · {mln(T['c1'])} млн ₸",
          "остаток сейчас" + (f" — запас ≈{T['stock_months']:.0f} мес. продаж" if T["stock_months"] else ""))]
    return k


def _fmt(v, kind):
    if v is None or v == "":
        return "—"
    if isinstance(v, str):
        return v
    if kind == "pct":
        return pct(v)
    if kind == "num":
        return "> 60" if v > 60 else f"{v:.0f}"
    return money(v)


# ------------------------------------------------------------------ Excel
def xlsx(R, who: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from . import xlsx_out
    wb = Workbook()
    dark = PatternFill("solid", fgColor="1F2328")
    thin = Side(style="thin", color="D0D7DE")
    ws = wb.active
    ws.title = "Итоги"
    ws["A1"] = "Анализ: " + _title(R)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Период {R['a']:%d.%m.%Y} – {R['b']:%d.%m.%Y}; остаток на {R['s_date']:%d.%m.%Y}"
                + (f"; организация {R['org']}" if R["org"] else ""))
    ws["A3"] = f"Сформировал: {who}, {xlsx_out.now_local():%d.%m.%Y %H:%M} · EDO.ANESCOM, данные 1С"
    r = 5
    for v, l in kpis(R):
        ws.cell(r, 1, v).font = Font(bold=True, size=12)
        ws.cell(r, 2, l)
        r += 1
    r += 1
    ws.cell(r, 1, "Выводы и что делать").font = Font(bold=True, size=12)
    r += 1
    for t in R["recs"]:
        ws.cell(r, 1, "• " + t).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.row_dimensions[r].height = 15 * (1 + len(t) // 110)
        r += 1
    r += 1
    for t in R["notes"]:
        ws.cell(r, 1, t).font = Font(italic=True, color="6B7280")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110

    def sheet(title, head, kinds, rows, tot):
        s = wb.create_sheet(title)
        s.append(head)
        for c in s[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = dark
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for row in rows + [tot]:
            s.append([None if (v is None) else v for v in row])
        n = s.max_row
        for i, k in enumerate(kinds, start=1):
            L = get_column_letter(i)
            s.column_dimensions[L].width = 48 if i == 1 else (22 if k == "text" else 13)
            for rr in range(2, n + 1):
                cell = s.cell(rr, i)
                cell.border = Border(bottom=thin)
                if isinstance(cell.value, (int, float)):
                    if k == "pct":
                        cell.value = cell.value / 100
                        cell.number_format = "0.0%"
                    else:
                        cell.number_format = "#,##0" if k in ("money", "int") else "0.0"
        for c in s[n]:
            c.font = Font(bold=True)
        s.row_dimensions[1].height = 42
        s.freeze_panes = "B2"
        s.auto_filter.ref = f"A1:{get_column_letter(len(head))}{n - 1}"
        return s

    sheet("По группам", *group_table(R))
    sheet("По товарам", *item_table(R))
    dead = [x for x in R["items"] if x["dead"]]
    sheet("Неликвид", ["Товар", "Группа", "Остаток, шт", "Сумма, ₸", "Последняя продажа"],
          ["text", "text", "int", "money", "text"],
          [[x["item"], x["group"], x["q1"], x["c1"], x["last_sale"] or "—"] for x in sorted(dead, key=lambda x: -x["c1"])],
          ["Итого", "", R["T"]["dead_q"], R["T"]["dead_c"], ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
def pdf(R, who: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    from xml.sax.saxutils import escape as X
    from . import pdfcert, xlsx_out
    pdfcert._ensure_font()
    F, FB = pdfcert._FONT, pdfcert._FONT_B
    RED, DARK, GRAY, LINE, GREEN = (colors.HexColor(x) for x in ("#C1272D", "#1F2328", "#6B7280", "#E5E7EB", "#2E8B57"))
    st = lambda size, bold=False, color=DARK, lead=None: ParagraphStyle(
        "s", fontName=FB if bold else F, fontSize=size, leading=lead or size * 1.3, textColor=color)
    from reportlab.lib.enums import TA_RIGHT
    H1, H2, P, SM, CELL, CELLB = st(20, True), st(13, True), st(9.5, lead=13), st(8, color=GRAY), st(7), st(7, True)
    CELLR, CELLBR = ParagraphStyle("r", parent=CELL, alignment=TA_RIGHT), ParagraphStyle("rb", parent=CELLB, alignment=TA_RIGHT)
    stamp = f"{xlsx_out.now_local():%d.%m.%Y %H:%M}"

    def deco(cv, doc):
        cv.saveState()
        w, h = landscape(A4)
        cv.setStrokeColor(RED)
        cv.setLineWidth(1.2)
        cv.line(14 * mm, h - 16 * mm, w - 14 * mm, h - 16 * mm)
        cv.setFont(FB, 11)
        cv.setFillColor(RED)
        cv.drawString(14 * mm, h - 13 * mm, "ANESCOM")
        cv.setFont(F, 8)
        cv.setFillColor(GRAY)
        cv.drawRightString(w - 14 * mm, h - 13 * mm, f"Товары и группы · {_title(R)[:80]}")
        cv.drawString(14 * mm, 8 * mm, f"EDO.ANESCOM · данные 1С · сформировал {who}, {stamp}")
        cv.drawRightString(w - 14 * mm, 8 * mm, str(doc.page))
        cv.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=21 * mm, bottomMargin=14 * mm, title=_title(R), author="EDO.ANESCOM")
    W = landscape(A4)[0] - 28 * mm
    E = []
    E.append(Paragraph(X(_title(R)), H1))
    E.append(Spacer(1, 4))
    E.append(Paragraph(f"Период {R['a']:%d.%m.%Y} – {R['b']:%d.%m.%Y} · остаток на {R['s_date']:%d.%m.%Y}"
                       + (f" · организация {X(R['org'])}" if R["org"] else "") + f" · позиций: {R['T']['n']}", P))
    E.append(Spacer(1, 8))
    cells, row = [], []
    for v, l in kpis(R):
        row.append([Paragraph(X(v), st(13, True)), Paragraph(X(l), st(8, color=GRAY, lead=10))])
        if len(row) == 3:
            cells.append(row)
            row = []
    kt = Table(cells, colWidths=[W / 3] * 3)
    kt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7F7F8")),
                            ("INNERGRID", (0, 0), (-1, -1), 4, colors.white),
                            ("LINEBEFORE", (0, 0), (0, -1), 3, RED), ("LINEBEFORE", (1, 0), (1, -1), 3, RED),
                            ("LINEBEFORE", (2, 0), (2, -1), 3, RED),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 8),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6)]))
    E.append(kt)
    E.append(Spacer(1, 10))
    E.append(Paragraph("Выводы и что делать", H2))
    E.append(Spacer(1, 3))
    for t in R["recs"] or ["Существенных проблем по отбору не найдено."]:
        E.append(Paragraph("• " + X(t), P))
        E.append(Spacer(1, 2))
    E.append(Spacer(1, 6))
    for t in R["notes"]:
        E.append(Paragraph(X(t), SM))

    def table(head, kinds, rows, tot, first=0.2):
        HS = st(6.6, True, colors.white, 8)
        HR = ParagraphStyle("hr", parent=HS, alignment=TA_RIGHT)
        data = [[Paragraph(X(h), HS if k == "text" else HR) for h, k in zip(head, kinds)]]
        jm = head.index("Маржа") if "Маржа" in head else -1
        for r in rows + [tot]:
            bold = r is tot
            line = []
            for j, (v, k) in enumerate(zip(r, kinds)):
                sty_ = (CELLB if bold else CELL) if k == "text" else (CELLBR if bold else CELLR)
                if j == jm and isinstance(v, (int, float)) and not bold:
                    sty_ = ParagraphStyle("m", parent=CELLBR, textColor=RED if v < 10 else GREEN if v >= 25 else DARK)
                line.append(Paragraph(X(str(v or "")) if k == "text" else X(_fmt(v, k)), sty_))
            data.append(line)
        wt = {"money": 1.35, "int": 0.85, "pct": 0.75, "num": 0.8}
        n_text = sum(1 for k in kinds if k == "text") - 1
        tot_w = sum(wt[k] for k in kinds if k != "text")
        unit = (1 - first - 0.09 * n_text) / tot_w
        ws = [W * (first if i == 0 else 0.09 if k == "text" else wt[k] * unit) for i, k in enumerate(kinds)]
        t = Table(data, colWidths=ws, repeatRows=1)
        sty = [("BACKGROUND", (0, 0), (-1, 0), DARK), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
               ("LINEBELOW", (0, 1), (-1, -1), 0.4, LINE), ("LINEABOVE", (0, -1), (-1, -1), 1.2, DARK),
               ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
               ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
        for i, k in enumerate(kinds):
            if k != "text":
                sty.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
        t.setStyle(TableStyle(sty))
        return t

    if len(R["groups"]) > 1:
        E.append(PageBreak())
        E.append(Paragraph("По группам товара", H2))
        E.append(Spacer(1, 4))
        E.append(table(*group_table(R), first=0.22))
    E.append(PageBreak())
    E.append(Paragraph("По товарам", H2))
    E.append(Paragraph("Цена поставщика — из документов «Приобретение товаров и услуг», без НДС. «С доставкой и пошлиной» — "
                       "себестоимость прихода в 1С (поступление + доп. расходы). Себестоимость продаж — из 1С; где месяц "
                       "не закрыт — оценка по средней цене. Запас = остаток сейчас ÷ средние продажи в месяц за период.", SM))
    E.append(Spacer(1, 4))
    E.append(table(*item_table(R, with_group=len(R["groups"]) > 1), first=0.17))
    doc.build(E, onFirstPage=deco, onLaterPages=deco)
    return buf.getvalue()
