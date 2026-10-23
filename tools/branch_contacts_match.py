"""Проект 19 (23.10.2026 22:10): сверка контактов из филиалов (Уральск — организации, Атырау — физлица)
с 1С (снимок an_partners) и Bitrix24 (живые данные crm.*.get).
Запуск без аргумента = отчёт (Б24 не меняется), `live` = запись в Б24: добавить недостающие телефоны/e-mail
филиала, адрес филиала (если в Б24 другой/пустой). Ничего не удаляет. Старые значения — branch_fix_live.jsonl.
Телефоны пишем в формате +7XXXXXXXXXX (правило пользователя 23.10 22:10).
Файлы филиалов: data\\reports\\filialy_uralsk_atyrau_kontragenty.xlsx, filial_atyrau_fizlica.xlsx.
Отчёт: data\\reports\\branch_contacts_<время>.xlsx (+ копия на рабочий стол)."""
import sys, os, re, json, time, collections, shutil
sys.path.insert(0, r"D:\anescom_edo")
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from app.analytics import conn
from app.b24people import _b24
from app.contactsync import norm_phone, norm_email, norm_name, same_name, split

LIVE = len(sys.argv) > 1 and sys.argv[1] == "live"
R = r"D:\anescom_edo\data\reports"
F1 = R + r"\filialy_uralsk_atyrau_kontragenty.xlsx"
F2 = R + r"\filial_atyrau_fizlica.xlsx"


def fmt(n):  # 7XXXXXXXXXX -> +7XXXXXXXXXX
    return "+" + n if n else ""


def phones_of(v):
    s = str(v or "").strip()
    if not s:
        return []
    if re.fullmatch(r"\d+(\.0)?", s):
        s = s.split(".")[0]
    out = []
    for part in re.split(r"[;,/\n]|\s{2,}| и ", s):
        d = re.sub(r"\D", "", part)
        if len(d) in (20, 21, 22):  # два номера слитно
            cands = [d[:len(d) // 2], d[len(d) // 2:]]
        else:
            cands = [part]
        for c in cands:
            n = norm_phone(c)
            if n and n not in out:
                out.append(n)
    return out


def emails_of(v):
    return [e for e in (norm_email(x) for x in re.split(r"[;,\s]+", str(v or ""))) if e]


# ---------- филиалы ----------
src = []  # (филиал, имя, регион, тел[], email[], адрес, лист)
wb = load_workbook(F1, data_only=True)
regions_only = []
for ws in wb.worksheets:
    rows = [r for r in ws.iter_rows(values_only=True) if any(r)]
    h = [str(x or "").strip() for x in rows[0]]
    for r in rows[1:]:
        d = dict(zip(h, r))
        nm = str(d.get("Наименование") or "").strip()
        if not nm:
            continue
        reg = str(d.get("Бизнес-регион") or "").strip()
        if "Номер телефона" in h:
            src.append(("Уральск", nm, reg, phones_of(d.get("Номер телефона")), emails_of(d.get("Эл.почта")),
                        str(d.get("Адрес") or "").strip(), ws.title, str(d.get("Номер телефона") or "")))
        else:
            regions_only.append((ws.title, nm, reg, str(d.get("Комментарий") or "").strip()))
wb2 = load_workbook(F2, data_only=True)
for r in wb2.worksheets[0].iter_rows(values_only=True):
    if not r or not r[0]:
        continue
    src.append(("Атырау", str(r[0]).strip(), str(r[1] or "").strip(), phones_of(r[2] if len(r) > 2 else ""), [], "",
                "Атырау физлица", str(r[2] if len(r) > 2 else "")))

# ---------- 1С и снимок Б24 ----------
with conn() as c:
    P = [dict(x) for x in c.execute("SELECT * FROM an_partners")]
    ct = [dict(x) for x in c.execute("SELECT id, name, origin FROM b24_snap_contact")]
    co = [dict(x) for x in c.execute("SELECT id, name, origin FROM b24_snap_company")]
partners = [p for p in P if p["rec"] == "partner" and not p["deleted"]]
by_nn = collections.defaultdict(list)
for p in partners:
    for n in (p["name"], p["full_name"]):
        if n:
            by_nn[norm_name(n)].append(p)
kids = collections.defaultdict(list)  # партнёр -> его контрагенты
for p in P:
    if p["rec"] == "contr" and p["partner_key"]:
        kids[p["partner_key"].lower()].append(p)
b24_by_origin = collections.defaultdict(list)
for kind, lst in (("company", co), ("contact", ct)):
    for x in lst:
        if x["origin"]:
            b24_by_origin[x["origin"].lower()].append((kind, x["id"], x["name"]))


def find_partner(nm, reg):
    c = by_nn.get(norm_name(nm)) or []
    if not c:
        c = [p for p in partners if same_name(nm, p["name"])]
    if len(c) > 1:
        c2 = [p for p in c if (p["region"] or "").strip().lower() == reg.lower()]
        c = c2 or c
        c3 = [p for p in c if p["sold"] or p["bought"]]
        c = c3 or c
    uniq = {p["key"].lower(): p for p in c}
    return list(uniq.values())


cache = {}


def b24get(kind, i):
    if (kind, i) not in cache:
        cache[(kind, i)] = _b24(f"crm.{kind}.get", {"id": i}) or {}
        time.sleep(0.2)
    return cache[(kind, i)]


rep, notfound, actions = [], [], []
log = open(R + r"\branch_fix_" + ("live" if LIVE else "dry") + ".jsonl", "a", encoding="utf-8")
st = collections.Counter()
for fil, nm, reg, ph, em, addr, sheet, raw in src:
    st["строк " + fil] += 1
    if not ph and raw.strip():
        st["нераспознанный номер"] += 1
    ps = find_partner(nm, reg)
    if not ps:
        notfound.append((fil, nm, reg, ", ".join(fmt(x) for x in ph) or raw, "; ".join(em), addr, "нет в 1С"))
        st["нет в 1С"] += 1
        continue
    if len(ps) > 1:
        st["несколько в 1С"] += 1
    p = ps[0]
    pk = p["key"].lower()
    keys = [pk] + [k["key"].lower() for k in kids.get(pk, [])]
    ph1 = []
    for r_ in [p] + kids.get(pk, []):
        ph1 += [n for n in (norm_phone(v) for v in split(r_["phones"])) if n and n not in ph1]
    em1 = []
    for r_ in [p] + kids.get(pk, []):
        em1 += [e for e in (norm_email(v) for v in split(r_["emails"])) if e and e not in em1]
    recs = []
    for k in keys:
        recs += b24_by_origin.get(k, [])
    org = fil == "Уральск"
    recs = [x for x in recs if (x[0] == "company") == org] or recs
    reg1c = (p["region"] or "").strip()
    if not recs:
        st["нет в Б24"] += 1
        notfound.append((fil, nm, reg, ", ".join(fmt(x) for x in ph), "; ".join(em), addr, f"нет в Б24 (1С: {p['name']})"))
        continue
    for kind, i, bname in recs:
        cur = b24get(kind, i)
        if not cur:
            continue
        ph2 = [norm_phone(v.get("VALUE", "")) or re.sub(r"\D", "", v.get("VALUE", "")) for v in cur.get("PHONE") or []]
        em2 = [(v.get("VALUE") or "").strip().lower() for v in cur.get("EMAIL") or []]
        a2 = (cur.get("ADDRESS") or "").strip()
        add_p = [x for x in ph if x not in ph2]
        add_e = [x for x in em if x not in em2]
        other_p = [x for x in ph2 if x and x not in ph and x not in ph1]
        set_a = bool(addr) and org and norm_name(addr) != norm_name(a2)
        no_1c_p = [x for x in ph if x not in ph1]
        no_1c_e = [x for x in em if x not in em1]
        if ph and not add_p and not add_e and not set_a:
            status = "совпадает"
        elif ph and set(ph) & set(ph2):
            status = "частично"
        elif ph2:
            status = "в Б24 другой номер"
        else:
            status = "в Б24 нет номера"
        st[status] += 1
        st["добавить тел."] += len(add_p); st["добавить e-mail"] += len(add_e); st["адрес"] += set_a
        st["тел. нет в 1С"] += len(no_1c_p)
        rep.append((fil, nm, p["name"], reg1c, "≠ " + reg if reg and reg1c and reg.lower() != reg1c.lower() else "",
                    f"{kind} {i}", bname,
                    ", ".join(fmt(x) for x in ph), ", ".join(fmt(x) for x in ph1), ", ".join(fmt(x) if len(x) == 11 else x for x in ph2),
                    ", ".join(fmt(x) for x in add_p), ", ".join(fmt(x) for x in other_p), ", ".join(fmt(x) for x in no_1c_p),
                    "; ".join(em), "; ".join(em1), "; ".join(em2), "; ".join(add_e), "; ".join(no_1c_e),
                    addr, a2, "заменить" if set_a else "", status))
        fl = {}
        if add_p:
            fl["PHONE"] = [{"VALUE": fmt(x), "VALUE_TYPE": "WORK"} for x in add_p]
        if add_e:
            fl["EMAIL"] = [{"VALUE": x, "VALUE_TYPE": "WORK"} for x in add_e]
        if set_a:
            fl["ADDRESS"] = addr[:250]
        if fl:
            log.write(json.dumps({"kind": kind, "id": i, "old": {"PHONE": cur.get("PHONE"), "EMAIL": cur.get("EMAIL"),
                                  "ADDRESS": cur.get("ADDRESS")}, "new": fl}, ensure_ascii=False) + "\n")
            if LIVE:
                _b24(f"crm.{kind}.update", {"id": i, "fields": fl, "params": {"REGISTER_SONET_EVENT": "N"}})
                time.sleep(0.3)
                st["записано"] += 1

# регионы (листы без контактов)
reg_rows = []
for sheet, nm, reg, com in regions_only:
    ps = find_partner(nm, reg)
    r1 = (ps[0]["region"] or "").strip() if ps else ""
    if not ps:
        reg_rows.append((sheet, nm, reg, "", "нет в 1С", com)); st["регионы: нет в 1С"] += 1
    elif r1.lower() != reg.lower():
        reg_rows.append((sheet, nm, reg, r1, "другой регион в 1С", com)); st["регионы: другой"] += 1
    else:
        st["регионы: совпадает"] += 1

out = Workbook()
ws = out.active; ws.title = "Сверка"
H = ["Филиал", "Наименование (филиал)", "Найдено в 1С", "Регион 1С", "Регион филиала", "Запись Б24", "Имя Б24",
     "Тел. филиал", "Тел. 1С", "Тел. Б24", "Добавить в Б24", "Прочие тел. в Б24 (нет у филиала и в 1С)",
     "Тел. филиала, которых нет в 1С", "E-mail филиал", "E-mail 1С", "E-mail Б24", "Добавить e-mail в Б24",
     "E-mail филиала, которых нет в 1С", "Адрес филиал", "Адрес Б24", "Адрес в Б24", "Итог"]
ws.append(H)
for r in rep:
    ws.append(list(r))
ws2 = out.create_sheet("Не найдено")
ws2.append(["Филиал", "Наименование (филиал)", "Регион", "Тел.", "E-mail", "Адрес", "Причина"])
for r in notfound:
    ws2.append(list(r))
ws3 = out.create_sheet("Регионы")
ws3.append(["Лист", "Наименование", "Регион филиала", "Регион 1С", "Итог", "Комментарий"])
for r in reg_rows:
    ws3.append(list(r))
for w in out.worksheets:
    for cl in w[1]:
        cl.font = Font(bold=True, color="FFFFFF"); cl.fill = PatternFill("solid", fgColor="2B4B5C")
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width = min(45, max(10, max(len(str(c.value or "")) for c in col[:60]) + 2))
    w.freeze_panes = "A2"
fn = R + time.strftime(r"\branch_contacts_%Y%m%d_%H%M.xlsx")
out.save(fn)
try:
    shutil.copy(fn, os.path.join(os.path.expanduser("~"), "Desktop", os.path.basename(fn)))
except Exception:
    pass
print("LIVE" if LIVE else "DRY", os.path.basename(fn))
print(json.dumps(dict(st), ensure_ascii=False))
