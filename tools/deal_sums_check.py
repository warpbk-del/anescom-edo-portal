"""Сверка сумм сделок Б24 с суммами заказов 1С (24.10.2026, проект 19). Только чтение.

Сделка Б24 ↔ заказ 1С: ORIGIN_ID сделки = GUID заказа (order_key в an_order_clients, сумма документа 1С).
Запуск на сервере:  D:\\anescom_edo\\.venv\\Scripts\\python.exe tools\\deal_sums_check.py
Итог: data\\reports\\deal_sums_YYYYMMDD_HHMM.xlsx и .json (сводка печатается).
"""
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.analytics import conn  # noqa: E402
from app.contactsync import _dump  # noqa: E402


def main():
    deals = _dump("deal", ["ID", "TITLE", "ORIGIN_ID", "ORIGINATOR_ID", "OPPORTUNITY", "CURRENCY_ID",
                           "STAGE_ID", "STAGE_SEMANTIC_ID", "CLOSED", "CATEGORY_ID", "DATE_CREATE"])
    with conn() as c:
        orders = {}
        for r in c.execute("SELECT order_key, order_no, date, org, sum FROM an_order_clients"):
            orders[r["order_key"]] = dict(r)
        span = c.execute("SELECT MIN(date) d1, MAX(date) d2, COUNT(*) n FROM an_order_clients").fetchone()
    rows, stat = [], Counter()
    for d in deals:
        o = orders.get((d.get("ORIGIN_ID") or "").strip())
        closed = d.get("CLOSED") == "Y"
        kind = "закрытая" if closed else "открытая"
        if not o:
            stat[f"{kind}: нет заказа в выгрузке 1С"] += 1
            continue
        b24 = float(d.get("OPPORTUNITY") or 0)
        one = float(o["sum"] or 0)
        diff = round(b24 - one, 2)
        if abs(diff) < 1:
            stat[f"{kind}: сумма совпадает"] += 1
            continue
        stat[f"{kind}: сумма РАЗНАЯ"] += 1
        rows.append({"deal": int(d["ID"]), "title": d.get("TITLE", ""), "closed": "да" if closed else "нет",
                     "stage": d.get("STAGE_ID", ""), "semantic": d.get("STAGE_SEMANTIC_ID", ""),
                     "order_no": o["order_no"], "order_date": (o["date"] or "")[:10], "org": o["org"],
                     "sum_b24": b24, "sum_1c": one, "diff": diff, "currency": d.get("CURRENCY_ID", "")})
    rows.sort(key=lambda r: (r["closed"] != "да", -abs(r["diff"])))
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    out = ROOT / "data" / "reports"
    summary = {"deals_b24": len(deals), "orders_1c": span["n"], "orders_period": [span["d1"], span["d2"]],
               "stat": dict(stat), "diff_rows": len(rows),
               "diff_closed_total": round(sum(r["diff"] for r in rows if r["closed"] == "да"), 2),
               "top": rows[:15]}
    (out / f"deal_sums_{ts}.json").write_text(json.dumps({"summary": summary, "rows": rows}, ensure_ascii=False),
                                              encoding="utf-8")
    try:
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Разные суммы"
        head = ["Сделка Б24", "Название", "Закрыта", "Стадия", "Заказ 1С", "Дата заказа", "Организация",
                "Сумма Б24", "Сумма 1С", "Разница (Б24−1С)", "Валюта"]
        ws.append(head)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r["deal"], r["title"], r["closed"], r["stage"], r["order_no"], r["order_date"], r["org"],
                       r["sum_b24"], r["sum_1c"], r["diff"], r["currency"]])
        ws2 = wb.create_sheet("Сводка")
        for k, v in summary["stat"].items():
            ws2.append([k, v])
        ws2.append(["Период заказов в выгрузке 1С", f"{span['d1']} — {span['d2']}"])
        wb.save(out / f"deal_sums_{ts}.xlsx")
    except Exception as e:  # noqa: BLE001
        summary["xlsx_error"] = str(e)
    print(json.dumps({k: v for k, v in summary.items() if k != "top"}, ensure_ascii=True))
    print("FILE", f"deal_sums_{ts}")


if __name__ == "__main__":
    main()
