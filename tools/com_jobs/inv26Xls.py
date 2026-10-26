# Excel для кладовщика: недостачи и излишки 2026 без пары (после 22-В), группы «есть с обеих сторон» — кандидаты на пересортицу
import re, openpyxl
from openpyxl.styles import Font, PatternFill
def grp(n): return ' '.join(re.sub(r'[\d,.]+(мм|м)?', '', n.lower()).split()[:2])
both = {grp(e['n']) for e in SS if e['r'] > 1e-6} & {grp(e['n']) for e in PP if e['r'] > 1e-6}
wb = openpyxl.Workbook(); H = ['Группа', 'Вид', 'Кол-во', 'Товар', 'Характеристика', 'Склад', 'Дата', 'Документ', 'Кандидат на пересортицу']
for title, L, vid in (('Недостачи', SS, 'списано недостачей'), ('Излишки', PP, 'оприходовано излишком')):
    ws = wb.active if title == 'Недостачи' else wb.create_sheet(title); ws.title = title; ws.append(H)
    for e in sorted([e for e in L if e['r'] > 1e-6], key=lambda e: (grp(e['n']), e['sk'], e['n'])):
        ws.append([grp(e['n']), vid, round(e['r'], 3), e['n'], e['x'], e['sk'], e['d'].strftime('%d.%m.%Y') if hasattr(e['d'], 'strftime') else str(e['d'])[:10], e['doc'], 'да' if grp(e['n']) in both else ''])
ws = wb.create_sheet('Кандидаты'); ws.append(H)
for e in sorted([e for e in SS + PP if e['r'] > 1e-6 and grp(e['n']) in both], key=lambda e: (grp(e['n']), e['sk'], e in PP)):
    ws.append([grp(e['n']), 'недостача' if e in SS else 'излишек', round(e['r'], 3), e['n'], e['x'], e['sk'], str(e['d'])[:10], e['doc'], 'да'])
for ws in wb.worksheets:
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='46B254')
    for col, wd in zip('ABCDEFGHI', (22, 20, 10, 55, 16, 26, 11, 48, 12)): ws.column_dimensions[col].width = wd
    ws.freeze_panes = 'A2'
fn = r'D:\anescom_edo\data\reports\Недостачи_излишки_2026_без_пары_ANESCOM_FIX.xlsx'; wb.save(fn); print('XLSX', fn, '| групп-кандидатов', len(both))
