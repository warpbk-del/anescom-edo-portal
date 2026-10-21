# Excel-смета открытия GREENTE из data.py: формулы, варианты Эконом/Оптимум/Максимум
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from data import *
wb=Workbook(); ws=wb.active; ws.title="Смета"
F="Arial"; teal="2B4B5C"; green="46B254"; ink="0B1F27"
fill=lambda c: PatternFill("solid",fgColor=c)
hdr=Font(name=F,bold=True,color="FFFFFF",size=10); bold=Font(name=F,bold=True,size=10); norm=Font(name=F,size=10); inp=Font(name=F,size=10,color="0000FF"); grey=Font(name=F,size=9,color="5B7079")
bd=Border(bottom=Side(style="thin",color="DCE6E4"))
ws["A1"]="GREENTE — смета открытия завода, "+EVENT; ws["A1"].font=Font(name=F,bold=True,size=14,color=teal)
ws["A2"]="Суммы в тенге. Синие — вводятся вручную, чёрные — формулы. КП — цена зафиксирована подрядчиком; Оценка — ориентир по рынку."; ws["A2"].font=grey
ws["A3"]="Гостей:"; ws["A3"].font=bold; ws["C3"]=GUESTS; ws["C3"].font=inp
ws["D3"]="Резерв:"; ws["D3"].font=bold; ws["E3"]=RESERVE; ws["E3"].font=inp; ws["E3"].number_format="0%"
cols=["Статья","Подрядчик","Кол-во","Цена (оптимум)","Сумма (оптимум)","Цена мин","Цена макс","Сумма эконом","Сумма максимум","Статус","Срок оплаты","Комментарий"]
r=5
for i,c in enumerate(cols,1):
    x=ws.cell(r,i,c); x.font=hdr; x.fill=fill(teal); x.alignment=Alignment(wrap_text=True,vertical="center")
r=6; subs=[]
for ci,(name,sub,items) in enumerate(CATS,1):
    ws.cell(r,1,f"{ci}. {name}").font=Font(name=F,bold=True,size=10,color=teal)
    for j in range(1,13): ws.cell(r,j).fill=fill("E8F5EA")
    r+=1; a=r
    for (n,who,q,p,lo,hi,st,due,cm) in items:
        ws.cell(r,1,n).font=norm; ws.cell(r,2,who).font=norm
        c=ws.cell(r,3,"=$C$3" if q==GUESTS and "Фуршет" in n else q); c.font=norm if isinstance(c.value,str) else inp
        ws.cell(r,4,p).font=inp; ws.cell(r,5,f"=C{r}*D{r}").font=norm
        ws.cell(r,6,lo).font=inp; ws.cell(r,7,hi).font=inp
        ws.cell(r,8,f"=C{r}*F{r}").font=norm; ws.cell(r,9,f"=C{r}*G{r}").font=norm
        ws.cell(r,10,st).font=Font(name=F,size=10,bold=True,color=green if st=="КП" else teal)
        ws.cell(r,11,due).font=grey; ws.cell(r,12,cm).font=grey
        for j in range(1,13): ws.cell(r,j).border=bd
        r+=1
    ws.cell(r,1,f"Итого: {name}").font=bold
    for col in "EHI": ws[f"{col}{r}"]=f"=SUM({col}{a}:{col}{r-1})"; ws[f"{col}{r}"].font=bold
    for j in range(1,13): ws.cell(r,j).fill=fill("F3F7F6")
    subs.append(r); r+=2
ws.cell(r,1,"Итого по статьям").font=bold
for col in "EHI": ws[f"{col}{r}"]="="+"+".join(f"{col}{s}" for s in subs); ws[f"{col}{r}"].font=bold
tot=r; r+=1
ws.cell(r,1,"Резерв").font=norm
for col in "EHI": ws[f"{col}{r}"]=f"={col}{tot}*$E$3"
res=r; r+=1
ws.cell(r,1,"ВСЕГО С РЕЗЕРВОМ").font=Font(name=F,bold=True,size=12,color="FFFFFF")
for col in "EHI": ws[f"{col}{r}"]=f"={col}{tot}+{col}{res}"; ws[f"{col}{r}"].font=Font(name=F,bold=True,size=12,color="FFFFFF")
for j in range(1,13): ws.cell(r,j).fill=fill(green)
grand=r; r+=1
ws.cell(r,1,"На одного гостя").font=norm
for col in "EHI": ws[f"{col}{r}"]=f"={col}{grand}/$C$3"
for row in ws.iter_rows(min_row=6,max_row=r):
    for c in row:
        if c.column in (4,5,6,7,8,9): c.number_format='# ##0'
widths=[46,20,8,13,15,11,11,14,15,10,20,50]
for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
ws.freeze_panes="B6"
d=wb.create_sheet("Уже сделано"); d.append(["Позиция","Статус","Факт, ₸"])
for c in d[1]: c.font=hdr; c.fill=fill(teal)
for a,s in DONE: d.append([a,s,None]); d.cell(d.max_row,3).font=inp
d.append(["Итого","",f"=SUM(C2:C{d.max_row})"]); d.cell(d.max_row,1).font=bold; d.cell(d.max_row,3).font=bold
d.column_dimensions["A"].width=40; d.column_dimensions["B"].width=20; d.column_dimensions["C"].width=15
s2=wb.create_sheet("Источники"); s2.append(["Источник","Ссылка"])
for c in s2[1]: c.font=hdr; c.fill=fill(teal)
for a,b in SOURCES: s2.append([a,b])
s2.column_dimensions["A"].width=80; s2.column_dimensions["B"].width=90
wb.save("/home/claude/GREENTE_byudzhet_otkrytiya.xlsx"); print("xlsx ok", grand)
